"""Audita el CSV sin modificarlo. Requiere openpyxl para leer los Excel."""
from pathlib import Path
from collections import defaultdict, Counter
import csv
import datetime as dt
import openpyxl

ROOT = Path(__file__).resolve().parent
ORIG = ROOT / 'Datos' / 'original'
SALIDA = ROOT / 'resultados'
ALIAS = {'series-historicas-precios-mayoristas.xlsx': 'series-historicas-precios-mayoristas_2015_2017.xlsx',
         'anex-SIPSA-SerieHistoricaMayorista-Dic2023(1).xlsx': 'anex-SIPSA-SerieHistoricaMayorista-2023.xlsx'}
normalizar = lambda s: str(s).casefold().replace('-', '').replace(' ', '')
with (ROOT / 'Datos/Base_panela_Corabastos_2015_2025.csv').open(encoding='utf-8-sig', newline='') as f:
    registros = list(csv.DictReader(f))
grupos = defaultdict(list)
for r in registros:
    grupos[r['archivo_fuente']].append(r)
resultados, notas = [], []
for nombre, entradas in grupos.items():
    local = ALIAS.get(nombre, nombre)
    try:
        libro = openpyxl.load_workbook(ORIG / local, read_only=True, data_only=True)
    except Exception as e:
        for r in entradas:
            resultados.append({**r, 'archivo_local': local, 'estado': 'no comprobable', 'detalle': str(e)})
        continue
    por_hoja = defaultdict(list)
    for r in entradas:
        nombres = [n for n in libro.sheetnames if normalizar(n) == normalizar(r['hoja_fuente'])]
        if len(nombres) != 1:
            resultados.append({**r, 'archivo_local': local, 'estado': 'no comprobable', 'detalle': 'Hoja no identificada inequívocamente'})
        else:
            por_hoja[nombres[0]].append(r)
    # También inspeccionar índices para recoger notas de fuente y unidad.
    for hoja_nombre in dict.fromkeys([*por_hoja, *[n for n in libro.sheetnames if normalizar(n) in ('indice', 'índice')]]):
        hoja = libro[hoja_nombre]
        requeridos = {int(r['fila_fuente']) for r in por_hoja[hoja_nombre]}
        filas = {}
        encabezado = None
        for numero, valores in enumerate(hoja.iter_rows(values_only=True), 1):
            if numero in requeridos:
                filas[numero] = valores
            if numero <= 15 or not isinstance(valores[0], (dt.datetime, dt.date)):
                textos = [str(x).strip() for x in valores if isinstance(x, str) and str(x).strip()]
                if textos and (numero <= 15 or any(t.lower().startswith(('fuente:', 'nota:', '*', 'dane', 'precio promedio')) for t in textos)):
                    notas.append(f'{local} | {hoja_nombre} | fila {numero} | ' + ' | '.join(textos))
            if numero <= 15 and any(str(v).strip().lower() == 'producto' for v in valores if v is not None):
                encabezado = [str(v).strip().lower() if v is not None else '' for v in valores]
        for r in por_hoja[hoja_nombre]:
            salida = {**r, 'archivo_local': local, 'hoja_local': hoja_nombre}
            fila = filas.get(int(r['fila_fuente']))
            if not fila or not encabezado:
                resultados.append({**salida, 'estado': 'no comprobable', 'detalle': 'Fila o encabezado no encontrado'}); continue
            def columna(opciones):
                return next((i for i, v in enumerate(encabezado) if any(x in v for x in opciones)), None)
            indices = [columna(['fecha']), columna(['producto']), columna(['mercado', 'fuente']), columna(['precio'])]
            if any(i is None for i in indices):
                resultados.append({**salida, 'estado': 'no comprobable', 'detalle': 'Faltan columnas identificables'}); continue
            fecha, producto, mercado, precio = (fila[i] for i in indices)
            checks = {'fecha': isinstance(fecha, (dt.date, dt.datetime)) and fecha.strftime('%d/%m/%Y') == r['fecha'],
                      'producto': producto == r['producto'], 'mercado': mercado == r['mercado'],
                      'precio': isinstance(precio, (int, float)) and precio == float(r['precio_cop_kg'])}
            por_kg = 'kilogramo' in encabezado[indices[3]]
            estado = 'difiere' if not all(checks.values()) else ('coincide' if por_kg else 'no comprobable')
            resultados.append({**salida, 'fecha_excel': str(fecha), 'producto_excel': producto,
                'mercado_excel': mercado, 'precio_excel': precio, 'encabezado_precio': encabezado[indices[3]],
                'estado': estado,
                'detalle': ', '.join(k for k, ok in checks.items() if not ok) or ('unidad por kg no explícita en encabezado' if not por_kg else ''),
                'unidad': 'por kilogramo; moneda COP no explícita' if por_kg else 'unidad y moneda no explícitas en encabezado'})
    libro.close()
campos = list(dict.fromkeys(k for r in resultados for k in r))
with (SALIDA / '15_verificacion_fuentes.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=campos); writer.writeheader(); writer.writerows(resultados)
(SALIDA / '15_notas_fuentes.txt').write_text('\n'.join(notas), encoding='utf-8')
print(Counter(r['estado'] for r in resultados))
